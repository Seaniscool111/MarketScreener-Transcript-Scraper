import os
import time
import re
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
from docx2pdf import convert # <--- New Library

# --- ADD YOUR COMPANIES HERE ---
COMPANIES = [
    {"name": "CISCO", "url": "https://www.marketscreener.com/quote/stock/CISCO-SYSTEMS-INC-4862/news-call-transcripts/"},
    {"name": "VEEV", "url": "https://www.marketscreener.com/quote/stock/VEEVA-SYSTEMS-INC-14691456/news-call-transcripts/"},
    {"name": "CRWD", "url": "https://www.marketscreener.com/quote/stock/CROWDSTRIKE-HOLDINGS-INC-59783691/news-call-transcripts/"}
]
# -------------------------------

BASE_FOLDER = "webscrape"

def sanitize_filename(filename):
    return re.sub(r'[\\/*?:"<>|]', "", filename).strip()

def save_transcript_as_docx(title, article_html, folder_path):
    """Saves transcript as a formatted Word doc."""
    safe_title = sanitize_filename(title)
    docx_path = os.path.join(folder_path, f"{safe_title}.docx")
    
    doc = Document()
    doc.add_heading(title, level=1)
    
    soup = BeautifulSoup(article_html, 'html.parser')
    
    # Extract headers, paragraphs, and speaker divs
    def is_target_tag(tag):
        if tag.name in ['p', 'h2', 'h3']:
            return True
        if tag.name == 'div' and tag.get('class'):
            # The speakers are typically in divs with a class like 'speech123456'
            return any('speech' in c for c in tag.get('class'))
        return False
        
    tags = soup.find_all(is_target_tag)
    
    last_text = ""
    for tag in tags:
        text = tag.get_text(" ", strip=True).replace('"', '').strip()
        if not text or text == last_text or len(text) < 5:
            continue
            
        p = doc.add_paragraph()
        
        is_speaker = tag.name == 'div' and tag.get('class') and any('speech' in c for c in tag.get('class'))
        
        # Speaker Detection: It's a speech div OR short lines without ending punctuation
        if is_speaker:
            name_span = tag.find('span', class_=lambda c: c and 'txt-bold' in c)
            # Find any other spans that are not the bold name span to get the title
            other_spans = tag.find_all('span', class_=lambda c: not c or 'txt-bold' not in c)
            title = " ".join(s.get_text(strip=True) for s in other_spans if s.get_text(strip=True))
            
            if name_span:
                name = name_span.get_text(strip=True)
                run_name = p.add_run(name)
                run_name.bold = True
                
                if title:
                    p.add_run("   ")
                    run_title = p.add_run(title)
                    run_title.font.size = Pt(9)
                    run_title.font.italic = True
            else:
                # Fallback if it's a speech div but lacks the expected spans
                run = p.add_run(text)
                run.bold = True
                
            p.paragraph_format.space_before = Pt(12) 
        elif len(text) < 45 and not text.endswith(('.', '?', '!')):
            run = p.add_run(text)
            run.bold = True
            p.paragraph_format.space_before = Pt(12) 
        else:
            p.add_run(text)
            p.paragraph_format.space_before = Pt(2)
            
        last_text = text

    # Save docx
    doc.save(docx_path)

def main():
    options = Options()
    options.add_experimental_option('excludeSwitches', ['enable-logging'])
    driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)

    # 1. Login ONCE
    print("\n--- STEP 1: LOGIN ---")
    driver.get("https://www.marketscreener.com/inscription/connexion.php")
    while True:
        if input("Type 'READY' when logged in: ").strip().upper() == "READY": break

    # 2. Loop through Companies
    for company in COMPANIES:
        print(f"\n--- PROCESSING: {company['name']} ---")
        comp_folder = os.path.join(BASE_FOLDER, company['name'])
        if not os.path.exists(comp_folder): os.makedirs(comp_folder)

        driver.get(company['url'])
        time.sleep(5)
        
        # Scroll to load all transcripts
        print("   Scrolling to load older transcripts...")
        last_height = driver.execute_script("return document.body.scrollHeight")
        while True:
            driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
            time.sleep(3) # Wait for content to load
            
            new_height = driver.execute_script("return document.body.scrollHeight")
            if new_height == last_height:
                # Wait a bit more just in case of slow network, and check again
                time.sleep(3)
                new_height = driver.execute_script("return document.body.scrollHeight")
                if new_height == last_height:
                    break
            last_height = new_height
        
        print("   Finished scrolling.")
        soup = BeautifulSoup(driver.page_source, 'html.parser')
        links = []
        for a in soup.find_all('a', href=True):
            href = a['href']
            if 'transcript' in href.lower() or 'transcript' in a.get_text().lower():
                full_url = "https://www.marketscreener.com" + href if href.startswith('/') else href
                if full_url not in links and company['url'] not in full_url:
                    links.append(full_url)

        print(f"Found {len(links)} transcripts.")

        # 3. Download and Save as DOCX
        for i, url in enumerate(links, 1):
            print(f"   [{i}/{len(links)}] Downloading...")
            driver.get(url)
            time.sleep(2) # Reduced from 4s to 2s
            
            inner_soup = BeautifulSoup(driver.page_source, 'html.parser')
            title = inner_soup.find('h1').get_text(strip=True) if inner_soup.find('h1') else f"Transcript_{i}"
            title = re.sub(r'^Transcript\s*[:\-]?\s*', '', title, flags=re.IGNORECASE)
            article_div = inner_soup.find('article') or inner_soup.find('div', id='newsContent')
            
            if article_div:
                save_transcript_as_docx(title, str(article_div), comp_folder)
                print(f"   ✓ Saved DOCX: {title[:40]}...")
            else:
                print(f"   ! Content not found.")
                
        # 4. Batch Convert DOCX to PDF (Much Faster!)
        print(f"\n   -> Batch converting all Word docs to PDF for {company['name']}...")
        try:
            convert(comp_folder)
        except Exception as e:
            print(f"   ! Batch PDF Conversion failed: {e}")
            
        # Clean up temporary DOCX files
        print(f"   -> Cleaning up temporary files...")
        for filename in os.listdir(comp_folder):
            if filename.endswith(".docx"):
                try:
                    os.remove(os.path.join(comp_folder, filename))
                except:
                    pass

    print("\n--- ALL TASKS FINISHED ---")
    driver.quit()

if __name__ == "__main__":
    main()